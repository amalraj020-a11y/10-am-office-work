"""Fixed artifact operations for Amala; no credentials, sending or model API calls.

Public work/<slug>/run.json fields: business_name, city, category, source_urls,
contact_source, template_name, template_url, license, model, commit, live_url,
screenshots, outreach_text, channel, send_status, message_id, updated_at, blocker.
Additional fields are ignored by the spreadsheet exporter.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FIELDS = ['business_name', 'city', 'category', 'source_urls', 'contact_source',
          'template_name', 'template_url', 'license', 'model', 'commit', 'live_url',
          'screenshots', 'outreach_text', 'channel', 'send_status', 'message_id',
          'updated_at', 'blocker']


def capture(slug, name):
    from playwright.sync_api import sync_playwright
    if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', slug):
        raise ValueError('Invalid business slug')
    url = f'https://amalraj020-a11y.github.io/10am-office-work/websites/{slug}/'
    out = ROOT / 'work/screenshots'
    out.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            page = browser.new_page(viewport={'width': 1440, 'height': 1000}, device_scale_factor=1)
            response = page.goto(url, wait_until='networkidle', timeout=60000)
            if not response or response.status != 200:
                raise RuntimeError('Live website is not available')
            text = page.locator('body').inner_text()
            if name.casefold() not in text.casefold() or 'Developed by MultiMax' not in text:
                raise RuntimeError('Wrong business or missing MultiMax footer')
            if re.search(r'lorem ipsum|your (?:company|business) name|123 street|testimonials? here', text, re.I):
                raise RuntimeError('Placeholder content remains')
            page.screenshot(path=str(out / f'{slug}-desktop.png'), full_page=True)
            page.set_viewport_size({'width': 390, 'height': 844})
            page.screenshot(path=str(out / f'{slug}-mobile.png'), full_page=True)
            if page.evaluate('document.documentElement.scrollWidth > innerWidth + 2'):
                raise RuntimeError('Mobile horizontal overflow')
        finally:
            browser.close()
    images = [out / f'{slug}-{view}.png' for view in ('desktop', 'mobile')]
    if len({hashlib.sha256(f.read_bytes()).hexdigest() for f in images}) != 2:
        raise RuntimeError('Screenshots must be distinct')
    print(json.dumps({'live_url': url, 'screenshots': [str(f.relative_to(ROOT)) for f in images]}))


def tracker():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook()
    sheet = book.active
    sheet.title = 'Business work'
    sheet.append(['record'] + FIELDS)
    for path in sorted((ROOT / 'work').glob('*/run.json')):
        if path.is_symlink():
            continue
        data = json.loads(path.read_text())
        row = [str(path.relative_to(ROOT))]
        for field in FIELDS:
            val = data.get(field, '')
            row.append(json.dumps(val, ensure_ascii=False) if isinstance(val, (list, dict)) else str(val))
        sheet.append(row)
    for row in sheet:
        for cell in row:
            # Force untrusted business/message strings to literal text, not Excel formulas.
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='164B43')
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 28
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    target = ROOT / 'work/office-tracker.xlsx'
    book.save(target)
    print(json.dumps({'tracker': str(target.relative_to(ROOT)), 'records': sheet.max_row - 1}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    commands = parser.add_subparsers(dest='command', required=True)
    cap = commands.add_parser('capture')
    cap.add_argument('--slug', required=True)
    cap.add_argument('--name', required=True)
    commands.add_parser('tracker')
    args = parser.parse_args()
    capture(args.slug, args.name) if args.command == 'capture' else tracker()
