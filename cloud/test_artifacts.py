import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest


class ArtifactTests(unittest.TestCase):
    def test_pages_preserve_distinct_shops_and_exclude_private_records(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'cloud').mkdir()
            shutil.copy2(Path(__file__).with_name('prepare_pages.py'), root / 'cloud/prepare_pages.py')
            for name in ('websites/sikkim-momo-gr-cafe', 'websites/another-shop', 'sites/older-shop'):
                path = root / name
                path.mkdir(parents=True)
                (path / 'index.html').write_text(name)
                (path / '.env').write_text('PRIVATE_TEST_VALUE')
            (root / 'work/screenshots').mkdir(parents=True)
            (root / 'work/screenshots/current.png').write_bytes(b'image-fixture')
            (root / 'work/private.json').write_text('PRIVATE_TEST_VALUE')
            subprocess.run([sys.executable, str(root / 'cloud/prepare_pages.py')], check=True)
            output = root / '_pages'
            self.assertEqual((output / 'index.html').read_text(), 'websites/sikkim-momo-gr-cafe')
            self.assertEqual((output / 'websites/another-shop/index.html').read_text(), 'websites/another-shop')
            self.assertTrue((output / 'sites/older-shop/index.html').exists())
            self.assertTrue((output / 'work/screenshots/current.png').exists())
            self.assertFalse((output / 'work/private.json').exists())
            self.assertFalse(list(output.rglob('.env')))

    def test_workbook_exports_only_public_fields_and_no_formulas(self):
        from openpyxl import load_workbook
        spec = importlib.util.spec_from_file_location('office_artifacts', Path(__file__).with_name('artifacts.py'))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            module.ROOT = root
            record = root / 'work/shop/run.json'
            record.parent.mkdir(parents=True)
            record.write_text(json.dumps({'business_name': '=1+1', 'send_status': 'not_sent', 'private_token': 'DO_NOT_EXPORT'}))
            module.tracker()
            book = load_workbook(root / 'work/office-tracker.xlsx', data_only=False)
            sheet = book.active
            self.assertEqual(sheet['B2'].value, '=1+1')
            self.assertEqual(sheet['B2'].data_type, 's')
            self.assertNotIn('DO_NOT_EXPORT', str(list(sheet.values)))
            self.assertEqual(sheet.max_row, 2)


if __name__ == '__main__':
    unittest.main()
